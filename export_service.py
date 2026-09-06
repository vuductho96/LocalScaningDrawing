import io
import csv
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExportService:
    """
    Dịch vụ tạo và định dạng báo cáo bóc tách kích thước & dung sai sang định dạng Excel (.xlsx) và CSV (.csv).
    """

    @staticmethod
    def generate_excel(
        drawing_name: str,
        global_constraints_summary: Optional[str],
        rows: List[Dict[str, Any]]
    ) -> io.BytesIO:
        """
        Tạo tệp Excel với định dạng chuyên nghiệp chuẩn kỹ thuật cơ khí.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DungSaiKichThuoc"
        ws.views.sheetView[0].showGridLines = True

        # Bộ font và màu sắc chuẩn
        title_font = Font(name="Segoe UI", size=15, bold=True, color="1E293B")
        subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        
        data_font = Font(name="Segoe UI", size=11, color="0F172A")
        nominal_font = Font(name="Segoe UI", size=11, bold=True, color="1E3A8A")
        global_font = Font(name="Segoe UI", size=10, italic=True, color="059669")
        local_font = Font(name="Segoe UI", size=10, color="2563EB")
        
        thin_border = Border(
            left=Side(style='thin', color="CBD5E1"),
            right=Side(style='thin', color="CBD5E1"),
            top=Side(style='thin', color="CBD5E1"),
            bottom=Side(style='thin', color="CBD5E1")
        )
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        # 1. Tiêu đề & Thông tin bản vẽ
        ws.merge_cells("A1:I1")
        ws["A1"] = f"BẢNG BÓC TÁCH KÍCH THƯỚC & DUNG SAI BẢN VẼ: {drawing_name}"
        ws["A1"].font = title_font
        ws["A1"].alignment = left_align
        ws.row_dimensions[1].height = 28

        ws.merge_cells("A2:I2")
        ws["A2"] = f"Quy tắc dung sai chung (Global Constraints): {global_constraints_summary or 'Theo số chữ số thập phân'}"
        ws["A2"].font = subtitle_font
        ws["A2"].alignment = left_align
        ws.row_dimensions[2].height = 20

        # 2. Tiêu đề các cột
        headers = [
            ("STT", 8, center_align),
            ("Số Lượng", 12, center_align),
            ("Ký Hiệu", 12, center_align),
            ("Nominal (Danh nghĩa)", 22, right_align),
            ("Dung Sai Trên (+)", 18, center_align),
            ("Dung Sai Dưới (-)", 18, center_align),
            ("Loại Dung Sai", 16, center_align),
            ("Kích Thước Đầy Đủ (Callout)", 32, left_align),
            ("Tọa Độ Crop (X, Y, W, H)", 26, center_align),
        ]

        header_row = 4
        ws.row_dimensions[header_row].height = 26
        for col_idx, (col_name, col_width, alignment) in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = alignment
            cell.border = thin_border
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = col_width

        # 3. Dữ liệu từng dòng kích thước
        row_start = 5
        for i, r in enumerate(rows):
            row_num = row_start + i
            ws.row_dimensions[row_num].height = 22
            
            # STT
            c1 = ws.cell(row=row_num, column=1, value=i + 1)
            c1.alignment = center_align
            c1.font = data_font
            c1.border = thin_border

            # Qty
            c2 = ws.cell(row=row_num, column=2, value=r.get("qty", ""))
            c2.alignment = center_align
            c2.font = data_font
            c2.border = thin_border

            # Prefix
            c3 = ws.cell(row=row_num, column=3, value=r.get("prefix", ""))
            c3.alignment = center_align
            c3.font = data_font
            c3.border = thin_border

            # Nominal (Nếu là góc độ DMS, giữ nguyên chuỗi đầy đủ)
            nom = r.get("nominal")
            nom_str = r.get("nominal_str", "")
            if r.get("tol_type") in ["angle", "angle_tol"] or any(c in str(nom_str) for c in ['°', "'", '"']):
                c4 = ws.cell(row=row_num, column=4, value=nom_str)
            else:
                c4 = ws.cell(row=row_num, column=4, value=nom if nom is not None else nom_str)
            c4.alignment = right_align
            c4.font = nominal_font
            c4.border = thin_border

            # Upper Tol
            c5 = ws.cell(row=row_num, column=5, value=r.get("upper_tol", ""))
            c5.alignment = center_align
            c5.font = data_font
            c5.border = thin_border

            # Lower Tol
            c6 = ws.cell(row=row_num, column=6, value=r.get("lower_tol", ""))
            c6.alignment = center_align
            c6.font = data_font
            c6.border = thin_border

            # Tol Type
            ttype = r.get("tol_type", "local")
            c7 = ws.cell(row=row_num, column=7, value=ttype.capitalize())
            c7.alignment = center_align
            c7.font = global_font if "global" in ttype else local_font
            c7.border = thin_border

            # Full Callout
            c8 = ws.cell(row=row_num, column=8, value=r.get("full_callout", ""))
            c8.alignment = left_align
            c8.font = data_font
            c8.border = thin_border

            # Tọa độ Crop (X, Y, W, H)
            b = r.get("box") or r.get("raw_box") or {}
            p_num = r.get("page", 0)
            if b and b.get("w", 0) > 0:
                coord_val = f"P{p_num + 1}: X={int(b.get('x', 0))}, Y={int(b.get('y', 0))}, W={int(b.get('w', 0))}, H={int(b.get('h', 0))}"
            else:
                coord_val = "-"
            c9 = ws.cell(row=row_num, column=9, value=coord_val)
            c9.alignment = center_align
            c9.font = data_font
            c9.border = thin_border

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    @staticmethod
    def generate_csv(rows: List[Dict[str, Any]]) -> io.StringIO:
        """
        Tạo nội dung file CSV.
        """
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "STT", "So Luong", "Ky Hieu", "Nominal", 
            "Dung Sai Tren (+)", "Dung Sai Duoi (-)", "Loai Dung Sai", 
            "Full Callout", "Toa Do Crop (X,Y,W,H)"
        ])
        for i, r in enumerate(rows):
            nom_val = r.get("nominal_str") if (
                r.get("tol_type") in ["angle", "angle_tol"] or 
                any(c in str(r.get("nominal_str", "")) for c in ['°', "'", '"'])
            ) else r.get("nominal", "")
            
            b = r.get("box") or r.get("raw_box") or {}
            p_num = r.get("page", 0)
            coord_val = f"P{p_num + 1}: X={int(b.get('x', 0))} Y={int(b.get('y', 0))} W={int(b.get('w', 0))} H={int(b.get('h', 0))}" if b and b.get("w", 0) > 0 else ""
            
            writer.writerow([
                i + 1,
                r.get("qty", ""),
                r.get("prefix", ""),
                nom_val,
                r.get("upper_tol", ""),
                r.get("lower_tol", ""),
                r.get("tol_type", ""),
                r.get("full_callout", ""),
                coord_val
            ])
        output.seek(0)
        return output
