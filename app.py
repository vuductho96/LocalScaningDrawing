import os
import uuid
import json
import csv
import io
import base64
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, FileResponse, StreamingResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel
from typing import Optional, List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from pdf_processor import PDFProcessor
from tolerance_parser import global_adaptive_learner
from ai_vision_service import global_ai_vision_service

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
STATIC_DIR = os.path.join(BASE_DIR, "static")
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")

os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(STATIC_DIR, exist_ok=True)
os.makedirs(TEMPLATES_DIR, exist_ok=True)

app = FastAPI(title="AutoScanText - OCR Ban Ve Ky Thuat", version="1.0.0")

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
templates = Jinja2Templates(directory=TEMPLATES_DIR)

pdf_processor = PDFProcessor(upload_dir=UPLOAD_DIR)

# Luu danh sach file theo id
active_files = {}

def get_or_restore_file(file_id: str) -> Optional[Dict[str, Any]]:
    if not file_id:
        return None
    if file_id in active_files:
        return active_files[file_id]

    # 1. Thu tim file sample mac dinh neu la sample_107 hoac sample
    if file_id in ["sample_107", "sample"]:
        sample_path = os.path.join(BASE_DIR, "107-M1457.pdf")
        if not os.path.exists(sample_path):
            sample_path = os.path.join(UPLOAD_DIR, "sample_drawing.pdf")
        if os.path.exists(sample_path):
            try:
                info = pdf_processor.get_pdf_info(sample_path)
                active_files[file_id] = {
                    "filename": os.path.basename(sample_path),
                    "path": sample_path,
                    "page_count": info["page_count"],
                    "pages": info["pages"]
                }
                return active_files[file_id]
            except Exception:
                pass

    # 2. Quet UPLOAD_DIR tim file da upload bat dau voi file_id_
    if os.path.exists(UPLOAD_DIR):
        for fname in os.listdir(UPLOAD_DIR):
            if fname.startswith(f"{file_id}_") and fname.lower().endswith(".pdf"):
                full_path = os.path.join(UPLOAD_DIR, fname)
                try:
                    info = pdf_processor.get_pdf_info(full_path)
                    orig_name = fname[len(file_id) + 1:]
                    active_files[file_id] = {
                        "filename": orig_name,
                        "path": full_path,
                        "page_count": info["page_count"],
                        "pages": info["pages"]
                    }
                    return active_files[file_id]
                except Exception:
                    pass

        # 3. Fallback: file co the luu nguyen ten hoac trung ten
        direct_path = os.path.join(UPLOAD_DIR, file_id)
        if os.path.exists(direct_path) and direct_path.lower().endswith(".pdf"):
            try:
                info = pdf_processor.get_pdf_info(direct_path)
                active_files[file_id] = {
                    "filename": file_id,
                    "path": direct_path,
                    "page_count": info["page_count"],
                    "pages": info["pages"]
                }
                return active_files[file_id]
            except Exception:
                pass

    return None

class CropRequest(BaseModel):
    file_id: str
    page_num: int = 0
    crop_box: Dict[str, float]  # { x, y, width, height } tu 0.0 den 1.0
    global_constraints: Optional[Dict[str, Any]] = None
    page_rotation: int = 0  # Goc xoay cua trang PDF: 0, 90, 180, 270
    crop_rotation: int = 0  # Goc xoay rieng cua vung crop: 0, 90, 180, 270

class ExportRequest(BaseModel):
    drawing_name: Optional[str] = "BanVeKyThuat"
    global_constraints_summary: Optional[str] = ""
    rows: List[Dict[str, Any]]

@app.get("/", response_class=FileResponse)
async def index(request: Request):
    html_path = os.path.join(TEMPLATES_DIR, "index.html")
    return FileResponse(html_path, media_type="text/html")

@app.get("/api/load-sample")
async def load_sample():
    sample_path = os.path.join(BASE_DIR, "107-M1457.pdf")
    if not os.path.exists(sample_path):
        sample_path = os.path.join(UPLOAD_DIR, "sample_drawing.pdf")
    
    file_id = "sample_107"
    filename = os.path.basename(sample_path)
    info = pdf_processor.get_pdf_info(sample_path)
    active_files[file_id] = {
        "filename": filename,
        "path": sample_path,
        "page_count": info["page_count"],
        "pages": info["pages"]
    }
    cache_path, w, h = pdf_processor.render_page(sample_path, page_num=0, dpi=200)
    return {
        "success": True,
        "file_id": file_id,
        "filename": filename,
        "page_count": info["page_count"],
        "page_width": w,
        "page_height": h,
        "initial_page_url": f"/api/page_image?file_id={file_id}&page=0"
    }

@app.post("/api/upload")
async def upload_pdf(file: UploadFile = File(...)):
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="Chi chap nhan file dinh dang .pdf")
    
    file_id = str(uuid.uuid4())[:8]
    save_name = f"{file_id}_{file.filename}"
    save_path = os.path.join(UPLOAD_DIR, save_name)

    with open(save_path, "wb") as f:
        content = await file.read()
        f.write(content)

    try:
        info = pdf_processor.get_pdf_info(save_path)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Loi doc file PDF: {str(e)}")

    active_files[file_id] = {
        "filename": file.filename,
        "path": save_path,
        "page_count": info["page_count"],
        "pages": info["pages"]
    }

    # Render truoc trang 0
    try:
        cache_path, w, h = pdf_processor.render_page(save_path, page_num=0, dpi=200)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Loi render trang PDF: {str(e)}")

    return {
        "success": True,
        "file_id": file_id,
        "filename": file.filename,
        "page_count": info["page_count"],
        "page_width": w,
        "page_height": h,
        "initial_page_url": f"/api/page_image?file_id={file_id}&page=0"
    }

@app.get("/api/page_image")
async def get_page_image(file_id: str, page: int = 0, rotation: int = 0):
    file_info = get_or_restore_file(file_id)
    if not file_info:
        raise HTTPException(status_code=404, detail="Khong tim thay file PDF da upload")
    
    pdf_path = file_info["path"]
    
    try:
        cache_path, w, h = pdf_processor.render_page(pdf_path, page_num=page, dpi=200, rotation=rotation)
        return FileResponse(cache_path, media_type="image/png")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Loi lay anh trang {page}: {str(e)}")

@app.post("/api/crop-ocr")
async def crop_ocr(req: CropRequest):
    file_info = get_or_restore_file(req.file_id)
    if not file_info:
        raise HTTPException(status_code=404, detail="Khong tim thay file PDF (vui long f5 hoac tai lai file)")

    pdf_path = file_info["path"]
    
    try:
        result = pdf_processor.crop_and_extract(
            pdf_path=pdf_path,
            page_num=req.page_num,
            crop_box=req.crop_box,
            global_constraints=req.global_constraints,
            dpi=200,
            page_rotation=req.page_rotation,
            crop_rotation=req.crop_rotation
        )
        return result
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Loi OCR: {str(e)}")

class CorrectionFeedbackRequest(BaseModel):
    raw_text: str
    corrected: Dict[str, Any]

@app.post("/api/feedback/correct")
async def save_correction_feedback(req: CorrectionFeedbackRequest):
    """
    Nhan du lieu chinh sua tu con nguoi va hoc thich ung lap tuc (1-Shot / Generalized).
    """
    try:
        res = global_adaptive_learner.learn_correction(req.raw_text, req.corrected)
        return res
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Loi luu phan hoi thich ung: {str(e)}")

@app.get("/api/feedback/rules")
async def get_adaptive_rules():
    """
    Lay danh sach cac quy tac da hoc tu nguoi dung de hien thi / quan ly tren UI.
    """
    return {
        "exact_matches": global_adaptive_learner.exact_matches,
        "generalized_rules": global_adaptive_learner.generalized_rules,
        "char_replacements": global_adaptive_learner.char_replacements
    }

@app.delete("/api/feedback/rules")
async def delete_adaptive_rule(key: str, rule_type: str = "exact"):
    """
    Xoa mot quy tac thich ung neu nguoi dung nhap nham.
    """
    success = global_adaptive_learner.delete_rule(key, rule_type)
    if not success:
        raise HTTPException(status_code=404, detail="Khong tim thay quy tac de xoa")
    return {"success": True, "message": f"Da xoa quy tac {key}"}

# ==============================================================================
# AI VISION ENDPOINTS (Gemini Multimodal Vision LLM)
# ==============================================================================

class AIVisionConfigRequest(BaseModel):
    api_key: str
    model_name: Optional[str] = "gemini-flash-latest"

class AIInspectCropRequest(BaseModel):
    file_id: str
    page_num: int = 0
    crop_box: Dict[str, float]
    raw_ocr_hint: Optional[str] = ""
    page_rotation: int = 0
    crop_rotation: int = 0

class AIAutoDetectRequest(BaseModel):
    file_id: str
    page_num: int = 0
    page_rotation: int = 0

@app.get("/api/ai-vision/status")
async def get_ai_vision_status():
    """Kiem tra trang thai cau hinh va san sang cua AI Vision."""
    return global_ai_vision_service.check_status()

@app.get("/api/ai-vision/usage")
async def get_ai_vision_usage():
    """Lay thong tin phan tram va so luong request / quota cua API Key."""
    return global_ai_vision_service.get_usage_stats()

@app.post("/api/ai-vision/config")
async def set_ai_vision_config(req: AIVisionConfigRequest):
    """Luu API key va cau hinh model AI Vision."""
    saved = global_ai_vision_service.save_config(req.api_key, req.model_name)
    status = global_ai_vision_service.check_status()
    return {"success": saved, "status": status}

@app.delete("/api/ai-vision/config")
async def clear_ai_vision_config():
    """Gỡ bỏ API key hoàn toàn khỏi hệ thống."""
    global_ai_vision_service.clear_config()
    status = global_ai_vision_service.check_status()
    return {"success": True, "status": status}

@app.post("/api/ai-vision/inspect-crop")
async def ai_inspect_crop(req: AIInspectCropRequest):
    """
    Dung AI Vision (Gemini) phan tich sau anh crop de giai ma kich thuoc kho,
    chu xoay nghieng, ky hieu GD&T hoac dung sai dac biet.
    """
    file_info = get_or_restore_file(req.file_id)
    if not file_info:
        raise HTTPException(status_code=404, detail="Khong tim thay file PDF")

    pdf_path = file_info["path"]
    try:
        crop_res = pdf_processor.crop_and_extract(
            pdf_path=pdf_path,
            page_num=req.page_num,
            crop_box=req.crop_box,
            page_rotation=req.page_rotation,
            crop_rotation=req.crop_rotation
        )
        
        thumb_data = crop_res.get("thumbnail", "")
        if "," in thumb_data:
            thumb_b64 = thumb_data.split(",")[1]
            img_bytes = base64.b64decode(thumb_b64)
        else:
            raise ValueError("Khong the trich xuat anh thumbnail")

        ai_res = global_ai_vision_service.inspect_crop(
            img_bytes, 
            raw_ocr_hint=req.raw_ocr_hint or crop_res.get("raw_text", "")
        )
        
        if not ai_res.get("success"):
            return {**crop_res, "ai_error": ai_res.get("error")}

        # Merge ket qua AI Vision
        merged = {
            **crop_res,
            "nominal": ai_res.get("nominal", crop_res.get("nominal")),
            "nominal_str": ai_res.get("nominal_str") or crop_res.get("nominal_str"),
            "upper_tol": ai_res.get("upper_tol", crop_res.get("upper_tol")),
            "lower_tol": ai_res.get("lower_tol", crop_res.get("lower_tol")),
            "qty": ai_res.get("qty", crop_res.get("qty")),
            "prefix": ai_res.get("prefix", crop_res.get("prefix")),
            "suffix": ai_res.get("suffix", crop_res.get("suffix")),
            "tol_type": ai_res.get("tol_type", crop_res.get("tol_type")),
            "full_callout": ai_res.get("full_callout", crop_res.get("full_callout")),
            "ai_explanation": ai_res.get("explanation", ""),
            "source": ai_res.get("source", "ai_vision")
        }
        return merged
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Loi AI Vision Inspect: {str(e)}")

@app.post("/api/ai-vision/auto-detect")
async def ai_auto_detect(req: AIAutoDetectRequest):
    """
    Dung AI Vision quet toan bo trang ban ve va phat hien tat ca cac cum kich thuoc.
    """
    file_info = get_or_restore_file(req.file_id)
    if not file_info:
        raise HTTPException(status_code=404, detail="Khong tim thay file PDF")

    pdf_path = file_info["path"]
    try:
        cache_path, w, h = pdf_processor.render_page(
            pdf_path, 
            page_num=req.page_num, 
            dpi=200, 
            rotation=req.page_rotation
        )
        with open(cache_path, "rb") as f:
            img_bytes = f.read()

        res = global_ai_vision_service.auto_detect_dimensions(
            img_bytes, 
            page_width=w, 
            page_height=h
        )
        return res
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Loi AI Auto-Detect: {str(e)}")

@app.post("/api/export-excel")
async def export_excel(req: ExportRequest):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DungSaiKichThuoc"
    ws.views.sheetView[0].showGridLines = True

    # Styles
    title_font = Font(name="Segoe UI", size=15, bold=True, color="1E293B")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    
    data_font = Font(name="Segoe UI", size=11, color="0F172A")
    nominal_font = Font(name="Segoe UI", size=11, bold=True, color="1E3A8A")
    global_font = Font(name="Segoe UI", size=10, italic=True, color="059669")
    local_font = Font(name="Segoe UI", size=10, color="2563EB")
    
    thin_border = Border(
        left=Side(style='thin', color="CBD5E1"),
        right=Side(style='thin', color="CBD5E1"),
        top=Side(style='thin', color="CBD5E1"),
        bottom=Side(style='thin', color="CBD5E1")
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    # Title & Metadata
    ws.merge_cells("A1:I1")
    ws["A1"] = f"BẢNG BÓC TÁCH KÍCH THƯỚC & DUNG SAI BẢN VẼ: {req.drawing_name}"
    ws["A1"].font = title_font
    ws["A1"].alignment = left_align
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:I2")
    ws["A2"] = f"Quy tắc dung sai chung (Global Constraints): {req.global_constraints_summary or 'Theo số chữ số thập phân'}"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = left_align
    ws.row_dimensions[2].height = 20

    # Headers
    headers = [
        ("STT", 8, center_align),
        ("Số Lượng", 12, center_align),
        ("Ký Hiệu", 12, center_align),
        ("Nominal (Danh nghĩa)", 22, right_align),
        ("Dung Sai Trên (+)", 18, center_align),
        ("Dung Sai Dưới (-)", 18, center_align),
        ("Loại Dung Sai", 16, center_align),
        ("Kích Thước Đầy Đủ (Callout)", 32, left_align),
        ("Tọa Độ Crop (X, Y, W, H)", 26, center_align),
    ]

    header_row = 4
    ws.row_dimensions[header_row].height = 26
    for col_idx, (col_name, col_width, alignment) in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = alignment
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

    # Rows Data
    row_start = 5
    for i, r in enumerate(req.rows):
        row_num = row_start + i
        ws.row_dimensions[row_num].height = 22
        
        # STT
        c1 = ws.cell(row=row_num, column=1, value=i + 1)
        c1.alignment = center_align
        c1.font = data_font
        c1.border = thin_border

        # Qty
        c2 = ws.cell(row=row_num, column=2, value=r.get("qty", ""))
        c2.alignment = center_align
        c2.font = data_font
        c2.border = thin_border

        # Prefix
        c3 = ws.cell(row=row_num, column=3, value=r.get("prefix", ""))
        c3.alignment = center_align
        c3.font = data_font
        c3.border = thin_border

        # Nominal
        nom = r.get("nominal")
        nom_str = r.get("nominal_str", "")
        # Neu la goc do (DMS), giu nguyen chuoi do phut giay day du (vi du: 0°10'36", 4°30'23")
        if r.get("tol_type") in ["angle", "angle_tol"] or any(c in str(nom_str) for c in ['°', "'", '"']):
            c4 = ws.cell(row=row_num, column=4, value=nom_str)
        else:
            c4 = ws.cell(row=row_num, column=4, value=nom if nom is not None else nom_str)
        c4.alignment = right_align
        c4.font = nominal_font
        c4.border = thin_border

        # Upper Tol
        c5 = ws.cell(row=row_num, column=5, value=r.get("upper_tol", ""))
        c5.alignment = center_align
        c5.font = data_font
        c5.border = thin_border

        # Lower Tol
        c6 = ws.cell(row=row_num, column=6, value=r.get("lower_tol", ""))
        c6.alignment = center_align
        c6.font = data_font
        c6.border = thin_border

        # Tol Type
        ttype = r.get("tol_type", "local")
        c7 = ws.cell(row=row_num, column=7, value=ttype.capitalize())
        c7.alignment = center_align
        c7.font = global_font if "global" in ttype else local_font
        c7.border = thin_border

        # Full Callout
        c8 = ws.cell(row=row_num, column=8, value=r.get("full_callout", ""))
        c8.alignment = left_align
        c8.font = data_font
        c8.border = thin_border

        # Tọa độ Crop (X, Y, W, H)
        b = r.get("box") or r.get("raw_box") or {}
        p_num = r.get("page", 0)
        if b and b.get("w", 0) > 0:
            coord_val = f"P{p_num + 1}: X={int(b.get('x',0))}, Y={int(b.get('y',0))}, W={int(b.get('w',0))}, H={int(b.get('h',0))}"
        else:
            coord_val = "-"
        c9 = ws.cell(row=row_num, column=9, value=coord_val)
        c9.alignment = center_align
        c9.font = data_font
        c9.border = thin_border

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"{req.drawing_name}_Tolerances.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@app.post("/api/export-csv")
async def export_csv(req: ExportRequest):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["STT", "So Luong", "Ky Hieu", "Nominal", "Dung Sai Tren (+)", "Dung Sai Duoi (-)", "Loai Dung Sai", "Full Callout", "Toa Do Crop (X,Y,W,H)"])
    for i, r in enumerate(req.rows):
        nom_val = r.get("nominal_str") if (r.get("tol_type") in ["angle", "angle_tol"] or any(c in str(r.get("nominal_str", "")) for c in ['°', "'", '"'])) else r.get("nominal", "")
        b = r.get("box") or r.get("raw_box") or {}
        p_num = r.get("page", 0)
        coord_val = f"P{p_num + 1}: X={int(b.get('x',0))} Y={int(b.get('y',0))} W={int(b.get('w',0))} H={int(b.get('h',0))}" if b and b.get("w", 0) > 0 else ""
        writer.writerow([
            i + 1,
            r.get("qty", ""),
            r.get("prefix", ""),
            nom_val,
            r.get("upper_tol", ""),
            r.get("lower_tol", ""),
            r.get("tol_type", ""),
            r.get("full_callout", ""),
            coord_val
        ])
    output.seek(0)
    filename = f"{req.drawing_name}_Tolerances.csv"
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode('utf-8-sig')),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

class CropCoordsRequest(BaseModel):
    file_id: str
    page_num: int = 0
    x: float
    y: float
    w: float
    h: float
    crop_rotation: int = 0
    global_constraints: Optional[Dict[str, Any]] = None

@app.post("/api/crop-by-coords")
async def crop_by_coords(req: CropCoordsRequest):
    """
    Endpoint ho tro AI / Script debug: Tu dong crop theo toa do pixel chinh xac tren ban ve.
    """
    file_info = get_or_restore_file(req.file_id)
    if not file_info:
        raise HTTPException(status_code=404, detail="Khong tim thay file PDF da upload")

    pdf_path = file_info["path"]
    cache_path, img_w, img_h = pdf_processor.render_page(pdf_path, page_num=req.page_num, dpi=200)

    norm_box = {
        "x": max(0.0, req.x / img_w),
        "y": max(0.0, req.y / img_h),
        "width": min(1.0, req.w / img_w),
        "height": min(1.0, req.h / img_h)
    }

    try:
        result = pdf_processor.crop_and_extract(
            pdf_path=pdf_path,
            page_num=req.page_num,
            crop_box=norm_box,
            global_constraints=req.global_constraints,
            dpi=200,
            page_rotation=0,
            crop_rotation=req.crop_rotation
        )
        return result
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Loi OCR theo toa do: {str(e)}")

if __name__ == "__main__":
    import uvicorn
    print("Khoi dong AutoScanText Web Server tai http://localhost:8000 ...")
    uvicorn.run("app:app", host="127.0.0.1", port=8000, reload=True)
